#!/usr/bin/env python3
"""Editable 5-minute deck. Reuses reference assets, never modifies reference.

python build_ppt.py --render
Optional PNG/JPG assets in assets/ replace the named image placeholders.
"""
from pathlib import Path
from io import BytesIO
import argparse
import json
import subprocess
import tempfile

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement

HERE = Path(__file__).resolve().parent
OUT = HERE / 'Internship_20260918.pptx'
REF = HERE / 'reference/Realtime_GS_Mapping.pptx'
NAVY, GRAY, LIGHT, ORANGE = '17136B', '595959', 'EEF0F7', 'D78039'
WHITE, LINE = 'FFFFFF', 'D4D6DF'
FONT = 'Arial'
TIMES = [10, 30, 25, 30, 25, 35, 35, 35, 50, 25]
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
ref = Presentation(REF)
slots = []
equations = []


def box(s, x, y, w, h, fill=WHITE, edge=None, shape=MSO_SHAPE.RECTANGLE):
    q = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    q.shadow.inherit = False
    q.fill.solid()
    q.fill.fore_color.rgb = RGBColor.from_string(fill)
    if edge:
        q.line.color.rgb = RGBColor.from_string(edge)
        q.line.width = Pt(1)
    else:
        q.line.fill.background()
    return q


def txt(s, x, y, w, h, text, size=17, color=GRAY, bold=False, center=False, font=FONT):
    q = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = q.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.12
        r = p.add_run()
        r.text = line
        r.font.name, r.font.size, r.font.bold = font, Pt(size), bold
        r.font.color.rgb = RGBColor.from_string(color)
    return q


def line(s, x1, y1, x2, y2, color=LINE, width=1.5):
    q = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    q.shadow.inherit = False
    q.line.color.rgb = RGBColor.from_string(color)
    q.line.width = Pt(width)
    return q


def arrow(s, x, y, w=.38, h=.20):
    return box(s, x, y, w, h, GRAY, shape=MSO_SHAPE.RIGHT_ARROW)


def slide(title=None, subtitle=None, source=None, notes=''):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    if title:
        box(s, 0, 0, 10, .70, NAVY)
        txt(s, .20, .075, 9.60, .53, title, 24, WHITE)
    if subtitle:
        txt(s, .40, .93, 9.20, .60, subtitle, 17)
    if source:
        txt(s, .28, 5.24, 9.05, .25, source, 7.8)
    txt(s, 9.43, 5.24, .30, .25, str(len(prs.slides)), 9, center=True)
    s.notes_slide.notes_text_frame.text = notes
    return s


def picture(s, blob, x, y, w, h):
    from PIL import Image
    im = Image.open(BytesIO(blob))
    ratio = min(w / im.width, h / im.height)
    pw, ph = im.width * ratio, im.height * ratio
    s.shapes.add_picture(BytesIO(blob), Inches(x + (w-pw)/2), Inches(y + (h-ph)/2), width=Inches(pw), height=Inches(ph))


def refpic(s, page, index, x, y, w, h):
    picture(s, ref.slides[page-1].shapes[index].image.blob, x, y, w, h)


def equation(s, key, latex, x, y, w, h, size=22):
    """High-resolution math; exact LaTeX retained for reproducible editing."""
    from matplotlib.mathtext import math_to_image
    from matplotlib.font_manager import FontProperties
    target=HERE/'equations'
    target.mkdir(exist_ok=True)
    path=target/f'{key}.png'
    math_to_image('$'+latex+'$',path,prop=FontProperties(size=size),dpi=360,color='#'+NAVY)
    picture(s,path.read_bytes(),x,y,w,h)
    equations.append({'slide':len(prs.slides),'key':key,'latex':latex})


def placeholder(s, key, label, x, y, w, h):
    """Image with contain fit, or an editable blank rectangle with a subtle label."""
    for ext in ('.png', '.jpg', '.jpeg'):
        path = HERE / 'assets' / (key + ext)
        if path.exists():
            picture(s, path.read_bytes(), x, y, w, h)
            return
    q = box(s, x, y, w, h, 'FCFCFE', LINE)
    q.name = 'IMAGE_PLACEHOLDER_' + key
    txt(s, x+.1, y+h/2-.15, w-.2, .3, label, 11, 'A0A2AD', center=True)
    slots.append({'slide':len(prs.slides), 'key':key, 'label':label, 'bounds_inches':[x,y,w,h]})


def takeaway(s, text):
    line(s, .4, 4.65, 9.6, 4.65)
    txt(s, .45, 4.76, 9.1, .34, text, 17, NAVY, center=True)


def labeled_box(s, x, y, w, h, label, fill=LIGHT, color=NAVY, size=16):
    box(s,x,y,w,h,fill)
    txt(s,x+.08,y+.04,w-.16,h-.08,label,size,color,center=True)


def build():
    s=slide(notes='00:00–00:10 | Presenter: Sangwon Park. Names: Sangwon Park, Dong Hwi Kim, Chaehyeon Song. Equal emphasis, as in the reference.')
    txt(s,.5,.70,9,1.25,'Real-time GS Mapping\non Aria',34,'111111',center=True)
    txt(s,.7,2.12,8.6,.6,'Building useful maps under limited compute',19,center=True)
    txt(s,.45,3.35,9.1,.52,'Sangwon Park, Dong Hwi Kim, Chaehyeon Song',22,center=True)
    txt(s,.7,3.98,8.6,.40,'RPM Lab, Seoul National University',13,center=True)
    txt(s,.7,4.70,8.6,.30,'Internship Presentation  ·  September 18, 2026',12,center=True)

    s=slide('Why Gaussian Splatting?','Represent a scene with colored, semi-transparent 3D Gaussians.',
            'Image source: supplied reference deck, slide 3. Kerbl et al., 3D Gaussian Splatting, SIGGRAPH 2023.',
            '00:10–00:40 | Representation → photorealism → fast rendering. Reference images are illustrative, not our results.')
    txt(s,.4,1.68,3.7,.35,'Point cloud',17,center=True)
    txt(s,4.75,1.68,4.85,.35,'Gaussian representation / rendering',17,center=True)
    refpic(s,3,4,.4,2.1,3.7,2.15)
    arrow(s,4.27,3.04)
    refpic(s,3,3,4.82,2.1,4.78,2.15)
    takeaway(s,'Photorealistic appearance + fast rendering')

    s=slide('Why Real-time Mapping?','A robot or AR user needs a useful map while moving.',
            'Image credit: Meta / Tom’s Guide (2025), as credited in the supplied reference deck.',
            '00:40–01:05 | Aria is the sensing platform; do not imply on-device computation. Map image slot is a real project output.')
    refpic(s,12,3,.45,1.85,3.65,2.1)
    txt(s,.55,4.03,3.45,.3,'Observe the environment',16,center=True)
    arrow(s,4.35,2.77,.48,.22)
    placeholder(s,'online_map','ONLINE MAP / DEMO FRAME',5.05,1.82,4.45,2.12)
    txt(s,5.1,4.03,4.35,.3,'Build and use the map now',16,center=True)
    takeaway(s,'Fast rendering does not guarantee fast map convergence.')

    s=slide('Mapping under Limited Compute','Tracking and mapping share a finite online compute budget.',
            'Conceptual allocation only; block widths are not measured timings.',
            '01:05–01:35 | Single shared compute resource; sequential input; no future views. Timing drawing is schematic.')
    txt(s,.45,1.77,1.5,.4,'Incoming frames',14)
    for i in range(8):
        labeled_box(s,2.05+i*.91,1.77,.65,.44,str(i+1),LIGHT,size=13)
    line(s,2.04,2.35,9.45,2.35,GRAY)
    txt(s,7.4,2.42,2,.30,'time →',12,center=True)
    txt(s,.45,2.93,1.45,.70,'Shared GPU\n(schematic)',14)
    for x,a,b in [(2.05,1.0,1.18),(4.45,1.36,.82),(6.85,.82,1.36)]:
        labeled_box(s,x,2.93,a,.65,'Track',NAVY,WHITE,13)
        labeled_box(s,x+a,2.93,b,.65,'Map',LIGHT,NAVY,13)
    txt(s,2.05,3.90,7.4,.38,'More mapping views compete for the same updates.',17)
    takeaway(s,'Use available observations within available computation.')

    s=slide('What Makes Online Mapping Difficult?','Two requirements: convincing appearance and reliable geometry.',
            notes='01:35–02:00 | Two distinct failure modes. Add actual crops; do not treat a conceptual image as experimental evidence.')
    txt(s,.45,1.7,4.3,.35,'Under-optimized appearance',18,NAVY,center=True)
    txt(s,5.2,1.7,4.3,.35,'Incorrect geometry',18,NAVY,center=True)
    placeholder(s,'underoptimized_crop','RECENT REGION / BLURRY RENDER',.55,2.19,4.1,1.65)
    placeholder(s,'floater_crop','FLOATER CLOSE-UP',5.3,2.19,4.1,1.65)
    txt(s,.55,3.98,4.1,.42,'Too little useful supervision',15,center=True)
    txt(s,5.3,3.98,4.1,.42,'Opacity in observed free space',15,center=True)
    takeaway(s,'How quickly can a newly observed region become reliable?')

    s=slide('When Should We Add More Views?','Use frames between keyframes, paced by completed optimization.',
            '§3.1  Optimization-Guided View Growth. Schematic; growth is capped by available observations.',
            '02:00–02:35 | N(t)-N0 approximately S(t)/kappa for available candidates and work-conserving admission. Exact cap: min(M(t),N0+floor(S(t)/kappa)). N is pool size; S is completed mapping updates, not GPU utilization. No hardware-invariance claim.')
    txt(s,.45,1.70,1.9,.50,'Tracking\nkeyframes',14)
    for i in range(9):
        key = i in (0,4,8)
        labeled_box(s,2.55+i*.74,1.73,.53,.53,'KF' if key else 'RGB',NAVY if key else LIGHT,WHITE if key else GRAY,11)
    txt(s,2.55,2.39,6.45,.30,'Intermediate RGB frames also carry mapping information.',13,center=True)
    labeled_box(s,.65,2.96,2.55,.62,'Completed updates',LIGHT,size=16)
    arrow(s,3.47,3.17,.48)
    labeled_box(s,4.20,2.96,2.2,.62,'Add a new view', 'FAEEDF',ORANGE,16)
    arrow(s,6.69,3.17,.48)
    labeled_box(s,7.39,2.96,1.96,.62,'Training pool',LIGHT,size=16)
    equation(s,'view_growth',r'N(t)-N_0\;\approx\;\frac{S(t)}{\kappa}',2.9,3.77,4.2,.48)
    txt(s,.55,4.35,8.9,.22,'N: training-pool size     S: completed GPU mapping updates     κ: updates per added view',11,center=True)
    takeaway(s,'Grow supervision at the pace of optimization.')

    s=slide('Which Views Should We Optimize?','Give under-trained observations more opportunities, while mixing views.',
            '§3.2  ERCB — one-step formulation. Normalization by candidate count N; illustrative counts.',
            '02:35–03:10 | Per-view count derivation from §3.2. Vnorm(n)=sum((n_i-mean(n))²)/(2N); Δ_i is its change after one selection. Minimize expected next-step variance change minus tau*entropy. Exact p_i=softmax(-beta*n_i), beta=1/(N*tau). N-normalization rescales temperature; it is not mean-normalization/CV². Candidate set fixed for this derivation; actual block sampling uses remaining candidates without replacement. No global equality/convergence guarantee. See MATH_NOTES.md.')
    txt(s,.5,1.64,4.0,.34,'Uneven optimization exposure',16,NAVY,center=True)
    base=2.92
    for i,(h,lbl) in enumerate([(.78,'A'),(.60,'B'),(.30,'C'),(.15,'D')]):
        x=.98+i*.83
        box(s,x,base-h,.46,h,NAVY if i<2 else ORANGE)
        txt(s,x-.05,base+.08,.56,.28,lbl,12,center=True)
    line(s,.70,base,4.18,base,GRAY)
    arrow(s,4.59,2.42,.52,.22)
    txt(s,5.35,1.64,4.05,.34,'Count-aware randomized selection',16,NAVY,center=True)
    for i,lbl in enumerate(['D','B','C','A']):
        labeled_box(s,5.62+i*.9,2.18,.63,.51,lbl,'FAEEDF' if lbl in ('D','C') else LIGHT,ORANGE if lbl in ('D','C') else NAVY,18)
    txt(s,5.25,2.81,4.35,.28,'Less-trained views + diverse sampling',13,center=True)
    line(s,.4,3.32,9.6,3.32)
    txt(s,.48,3.52,1.85,.48,'Normalized variance\n+ entropy',12,NAVY)
    equation(s,'ercb_objective',r'\min_{p\in\Delta_N}\;\mathbb{E}_{i\sim p}\!\left[\Delta_i V_{\mathrm{norm}}(\mathbf{n})\right]-\tau H(p)',2.45,3.46,6.95,.58)
    txt(s,.48,4.19,1.85,.33,'Softmax solution',12,NAVY)
    equation(s,'ercb_softmax',r'p_i^*=\operatorname{softmax}_i(-\beta n_i)=\frac{e^{-\beta n_i}}{\sum_j e^{-\beta n_j}}',2.45,4.08,6.95,.57)
    equation(s,'ercb_variance',r'V_{\mathrm{norm}}(\mathbf{n})=\frac{1}{2N}\sum_i(n_i-\bar n)^2',.55,4.83,4.3,.29,size=15)
    txt(s,5.05,4.81,4.40,.35,'nᵢ: selection count    H: sampling entropy',11,center=True)

    s=slide('Keeping Observed Free Space Empty','A good color match can still hide an incorrect 3D structure.',
            'Carve concept from the current manuscript. Schematic; integrated geometry results remain to be validated.',
            '03:10–03:45 | One-sided Carve intuition: suppress opacity before a depth-supported surface; leave unobserved space unconstrained. Carve/Hit choice still under study.')
    box(s,2.03,2.18,4.52,1.73,'F4F5FA')
    txt(s,2.1,1.75,4.3,.33,'Observed free space',16,NAVY,center=True)
    labeled_box(s,.5,2.73,1.18,.62,'Camera',LIGHT,size=14)
    line(s,1.69,3.04,8.8,3.04,GRAY,2)
    box(s,6.58,2.13,.12,1.85,NAVY)
    txt(s,6.16,4.03,1.03,.45,'Observed\nsurface',12,center=True)
    box(s,3.70,2.78,.72,.53,'EFC49E',ORANGE,MSO_SHAPE.OVAL)
    txt(s,3.17,3.45,1.8,.35,'Reduce opacity',14,ORANGE,center=True)
    txt(s,7.22,2.19,2.05,.65,'Unobserved\nspace',15,center=True)
    txt(s,7.17,3.42,2.15,.47,'No free-space claim',12,center=True)
    takeaway(s,'Use causal depth evidence to suppress premature opacity.')

    s=slide('Current Progress & Evaluation','View growth and scheduling are implemented; joint validation is ongoing.',
            'Ongoing work: evaluate appearance, free-space consistency, and surface completeness together.',
            '03:45–04:35 | Fill comparison and metric before presentation. No invented numbers. Report dataset, GPU, budget, held-out split, repeats, tail policy. Identify personal contribution after confirming division of work.')
    txt(s,.45,1.69,4.3,.33,'Baseline',17,NAVY,center=True)
    txt(s,5.2,1.69,4.3,.33,'Proposed setting',17,NAVY,center=True)
    placeholder(s,'result_baseline','MATCHED VIEW / BASELINE',.55,2.13,4.1,1.73)
    placeholder(s,'result_proposed','MATCHED VIEW / PROPOSED',5.3,2.13,4.1,1.73)
    placeholder(s,'result_metric','KEY METRIC / EVALUATION CONDITIONS',.55,4.08,8.85,.48)
    txt(s,.55,4.77,8.85,.34,'Next: matched-budget appearance + geometry evaluation',16,NAVY,center=True)

    s=slide('Takeaway & Next Step',notes='04:35–05:00 | Recap research direction, not an unverified success claim. Hold conclusion for final five seconds.')
    txt(s,.60,1.05,8.8,.9,'Make each unit of online computation\nproduce a more useful map.',26,NAVY,center=True)
    for x,n,head,body in [(.55,'01','When to add','Pace view growth'),(3.75,'02','What to optimize','Balance exposure'),(6.95,'03','Where to constrain','Use free-space evidence')]:
        txt(s,x,2.43,2.5,.40,n,24,ORANGE,center=True)
        txt(s,x,3.04,2.5,.40,head,18,NAVY,center=True)
        txt(s,x-.1,3.63,2.7,.42,body,14,center=True)
    takeaway(s,'Next: validate appearance and geometry gains at the native input rate.')

    assert len(prs.slides)==10 and sum(TIMES)==300
    for i,s in enumerate(prs.slides):
        tr=OxmlElement('p:transition')
        tr.set('advClick','1')
        tr.set('advTm',str(TIMES[i]*1000))
        s._element.append(tr)
        for sh in s.shapes:
            # LibreOffice also honors the theme effectRef even with empty effectLst.
            for effect in sh._element.xpath('.//a:effectRef'):
                effect.set('idx','0')
            assert sh.left>=0 and sh.top>=0, (i+1,sh.name)
            assert sh.left+sh.width <= prs.slide_width+10, (i+1,sh.name,'right')
            assert sh.top+sh.height <= prs.slide_height+10, (i+1,sh.name,'bottom')
    prs.core_properties.title='Real-time GS Mapping on Aria'
    prs.core_properties.author='Sangwon Park, Dong Hwi Kim, Chaehyeon Song'
    prs.core_properties.subject='Five-minute internship presentation, 2026-09-18'
    prs.save(OUT)
    (HERE/'image_slots.json').write_text(json.dumps(slots,indent=2),encoding='utf-8')
    (HERE/'equations/equations.json').write_text(json.dumps(equations,indent=2),encoding='utf-8')
    print(OUT)


def render():
    from PIL import Image, ImageDraw
    target=HERE/'rendered'
    target.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='intern-lo-') as profile:
        subprocess.run(['libreoffice','-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf','--outdir',str(HERE),str(OUT)],check=True,timeout=120)
    pdf=OUT.with_suffix('.pdf')
    subprocess.run(['pdftoppm','-r','144','-png',str(pdf),str(target/'slide')],check=True,timeout=120)
    page_text=subprocess.check_output(['pdftotext',str(pdf),'-'],text=True).split('\f')
    assert len(page_text)-1==10
    thumbs=[]
    report=[]
    for i in range(10):
        path=target/f'slide-{i+1:02d}.png'
        im=Image.open(path).convert('RGB')
        im.thumbnail((640,360))
        panel=Image.new('RGB',(660,400),'#e9e9ee')
        panel.paste(im,((660-im.width)//2,10))
        ImageDraw.Draw(panel).text((12,378),f'{i+1:02d}  |  {TIMES[i]} sec',fill='black')
        thumbs.append(panel)
        report.append({'slide':i+1,'text_chars':len(page_text[i]),'render':str(path.relative_to(HERE))})
    contact=Image.new('RGB',(1320,2000),'#e9e9ee')
    for i,im in enumerate(thumbs):
        contact.paste(im,((i%2)*660,(i//2)*400))
    contact.save(target/'contact_sheet.png')
    (target/'validation.json').write_text(json.dumps({'slides':10,'duration_seconds':sum(TIMES),'shape_bounds':'pass','pages':report},indent=2))
    print(OUT.with_suffix('.pdf'))
    print(target/'contact_sheet.png')


if __name__=='__main__':
    ap=argparse.ArgumentParser()
    ap.add_argument('--render',action='store_true')
    args=ap.parse_args()
    build()
    if args.render:
        render()
